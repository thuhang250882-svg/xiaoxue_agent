import tempfile
import unittest
import zipfile
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook

from revise_office import revise


class ReviseOfficeTest(unittest.TestCase):
    def decision(self, match="Old term", replacement="New term"):
        return [{"id": "R-001", "match_text": match, "replacement": replacement, "comment": "需要更新"}]

    def test_docx_outputs_annotation_and_clean_replacement(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.docx"
            annotated = root / "annotated.docx"
            final = root / "final.docx"
            document = Document()
            document.add_paragraph("Prefix Old term suffix")
            document.save(source)

            result = revise(source, annotated, final, self.decision())

            self.assertEqual(len(result["annotated"]["applied"]), 1)
            self.assertEqual(list(Document(annotated).comments)[0].author, "AI审核")
            self.assertEqual("".join(p.text for p in Document(final).paragraphs), "Prefix New term suffix")

    def test_xlsx_outputs_cell_comment_and_replacement(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.xlsx"
            annotated = root / "annotated.xlsx"
            final = root / "final.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "Prefix Old term suffix"
            workbook.save(source)

            revise(source, annotated, final, self.decision())

            marked = load_workbook(annotated)
            clean = load_workbook(final)
            self.assertIn("AI审核", marked.active["A1"].value)
            self.assertEqual(clean.active["A1"].value, "Prefix New term suffix")
            marked.close()
            clean.close()

    def test_pptx_outputs_visible_marker_and_replacement(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.pptx"
            annotated = root / "annotated.pptx"
            final = root / "final.pptx"
            p_ns = "http:" + "//schemas.openxmlformats.org/presentationml/2006/main"
            a_ns = "http:" + "//schemas.openxmlformats.org/drawingml/2006/main"
            slide = f'<p:sld xmlns:p="{p_ns}" xmlns:a="{a_ns}"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Prefix Old term suffix</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>'.encode()
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("ppt/slides/slide1.xml", slide)

            revise(source, annotated, final, self.decision())

            with zipfile.ZipFile(annotated) as archive:
                self.assertIn("AI", archive.read("ppt/slides/slide1.xml").decode("utf-8"))
            with zipfile.ZipFile(final) as archive:
                text = archive.read("ppt/slides/slide1.xml").decode("utf-8")
                self.assertIn("New term", text)
                self.assertNotIn("Old term", text)

    def test_pdf_outputs_highlight_and_replacement(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.pdf"
            annotated = root / "annotated.pdf"
            final = root / "final.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "付款期限", fontname="china-s")
            document.save(source)
            document.close()

            revise(source, annotated, final, self.decision("付款期限", "付款时间"))

            marked = fitz.open(annotated)
            clean = fitz.open(final)
            self.assertIsNotNone(marked[0].first_annot)
            self.assertIn("付款时间", "".join(page.get_text() for page in clean))
            marked.close()
            clean.close()


if __name__ == "__main__":
    unittest.main()
