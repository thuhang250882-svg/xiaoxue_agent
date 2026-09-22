# -*- coding: utf-8 -*-
"""
周报文档生成器 — 六套格式统一入口

一套代码覆盖六种报送口径，按格式名选择排版参数：

    格式名           用途
    ─────────────── ──────────────────────────────────────────
    甲方周报          塔里木油田周报（页边距 2.5/2.5/2.4/2.4，正文仿宋16pt）
    甲方旬报          塔里木油田旬报（同上，结构为服务动态/专项治理/问题）
    联席会汇报        分院周工作联席会五段式公文（页边距 3.7/3.5/2.8/2.6，全文不加粗）
    分院党建周报      党建+信息化两块（正文仿宋14pt，最短）
    信息化周报        项目部内部（正文仿宋16pt，含A12变体）
    个人周报          个人留档（默认字体，两段式）

用法：

    from report_builder import ReportBuilder
    rb = ReportBuilder('联席会汇报')
    rb.build([
        ('title', '塔里木分院'),
        ('time',  '（2026-08-29 ~ 2026-09-11）'),
        ('h1',    '一、周、月度会安排工作落实情况'),
        ('p',     '本周例会安排事项已按要求推进落实。'),
        ('sign',  '汇报人：胡向峰'),
    ])
    rb.save('out.docx')

块类型：title/time/h1/h2/lead/p/sign/table/table_header/table_widths/row/table_end
    title 大标题      h1 一级标题      h2 二级标题
    lead  板块引导行（如"下周重点工作"）
    p     正文        sign 署名行
    table 第一行表头（分号分隔）  table_header 第二行表头
    table_widths 列宽权重          row 表格数据行
    table_end 结束当前表格

命令行：
    python report_builder.py --list
    python report_builder.py --selfcheck
    python report_builder.py --format 联席会汇报 --blocks blocks.txt --out out.docx
"""
import argparse
import logging
import os
import re
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

logger = logging.getLogger(__name__)

FONT_TITLE = '方正小标宋简体'
FONT_HEI = '方正黑体简体'
FONT_KAI = '方正楷体简体'
FONT_FANG = '方正仿宋简体'

_H1_HEI = {'font': FONT_HEI, 'size': 16, 'bold': True, 'indent': 2}
_H1_HEI_PLAIN = {'font': FONT_HEI, 'size': 16, 'bold': False, 'indent': 2}
_P_FANG16 = {'font': FONT_FANG, 'size': 16, 'bold': False, 'indent': 2}
_H1_HEI_PLAIN_EXACT = {**_H1_HEI_PLAIN, 'line_exact': 28}
_P_FANG16_EXACT = {**_P_FANG16, 'line_exact': 28}
_TITLE_SONG = {'font': FONT_TITLE, 'size': 22, 'bold': False, 'indent': 0}
_MARGIN_A = (2.5, 2.5, 2.4, 2.4)

FORMATS = {
    '甲方周报': {
        'label': '钻试修井场信息化服务周工作情况汇报（塔里木油田）',
        'margins': _MARGIN_A,
        'normal': (FONT_FANG, 16),
        'styles': {
            'title': _TITLE_SONG,
            'h1': _H1_HEI,
            'p': _P_FANG16,
        },
    },
    '甲方旬报': {
        'label': '钻试修井场信息化服务旬报（塔里木油田）',
        'margins': _MARGIN_A,
        'normal': (FONT_FANG, 16),
        'styles': {
            'title': _TITLE_SONG,
            'h1': _H1_HEI,
            'p': _P_FANG16,
        },
    },
    '联席会汇报': {
        'label': '分院周工作联席会汇报（五段式公文）',
        'margins': (3.7, 3.5, 2.8, 2.6),
        'normal': (FONT_FANG, 16),
        'styles': {
            'title': {'font': FONT_TITLE, 'size': 22, 'bold': False,
                      'indent': 7, 'line_exact': 32},
            'time': _P_FANG16_EXACT,
            'h1': _H1_HEI_PLAIN_EXACT,
            'h2': {'font': FONT_KAI, 'size': 16, 'bold': False, 'indent': 2, 'line_exact': 28},
            'p': _P_FANG16_EXACT,
            'sign': {'font': FONT_FANG, 'size': 16, 'bold': False, 'indent': 0, 'line_exact': 28},
        },
    },
    '分院党建周报': {
        'label': '党建工作及信息化服务周工作（分院内部）',
        'margins': _MARGIN_A,
        'normal': (FONT_FANG, 14),
        'styles': {
            'title': _TITLE_SONG,
            'h1': _H1_HEI,
            'lead': {'font': FONT_FANG, 'size': 14, 'bold': True, 'indent': 2},
            'p': {'font': FONT_FANG, 'size': 14, 'bold': False, 'indent': 2},
        },
    },
    '信息化周报': {
        'label': '西部钻探信息化服务项目部周工作汇报（内部）',
        'margins': _MARGIN_A,
        'normal': (FONT_FANG, 16),
        'styles': {
            'title': _TITLE_SONG,
            'h1': _H1_HEI,
            'p': _P_FANG16,
        },
    },
    '个人周报': {
        'label': '个人工作周报（默认字体，两段式）',
        'margins': (2.54, 2.54, 3.17, 3.17),
        'normal': None,
        'styles': {
            'lead': {'font': None, 'size': None, 'bold': False, 'indent': 0},
            'p': {'font': None, 'size': None, 'bold': False, 'indent': 0},
        },
    },
}


class ReportBuilder:
    """按格式名构建周报 docx"""

    def __init__(self, fmt_name):
        if fmt_name not in FORMATS:
            raise ValueError(
                f'未知格式：{fmt_name}。可选：{"、".join(FORMATS)}')
        self.fmt_name = fmt_name
        self.spec = FORMATS[fmt_name]
        self.doc = None
        self._table_spec = None

    # ---------- 底层排版 ----------

    def _setup_page(self):
        sec = self.doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        t, b, l, r = self.spec['margins']
        sec.top_margin, sec.bottom_margin = Cm(t), Cm(b)
        sec.left_margin, sec.right_margin = Cm(l), Cm(r)

    def _setup_normal(self):
        normal_cfg = self.spec.get('normal')
        if not normal_cfg:
            return
        name, size = normal_cfg
        normal = self.doc.styles['Normal']
        normal.font.name = name
        normal.font.size = Pt(size)
        rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
            rfonts.set(qn(attr), name)

    def _set_run(self, run, cfg):
        if cfg.get('font'):
            run.font.name = cfg['font']
            rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
                rfonts.set(qn(attr), cfg['font'])
        if cfg.get('size'):
            run.font.size = Pt(cfg['size'])
        run.font.bold = bool(cfg.get('bold'))

    @staticmethod
    def _set_indent_chars(paragraph, chars):
        """按"字符"为单位设置首行缩进，随字号自适应"""
        ppr = paragraph._element.get_or_add_pPr()
        ind = ppr.find(qn('w:ind'))
        if ind is None:
            ind = ppr.makeelement(qn('w:ind'), {})
            ppr.append(ind)
        ind.set(qn('w:firstLineChars'), str(int(chars * 100)))
        ind.set(qn('w:firstLine'), str(int(chars * 160)))

    def _add_text(self, text, cfg):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        self._set_run(run, cfg)
        if cfg.get('indent'):
            self._set_indent_chars(p, cfg['indent'])
        if cfg.get('line_exact'):
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(cfg['line_exact'])
        return p

    # ---------- 对外接口 ----------

    def build(self, blocks):
        """
        Args:
            blocks: [(kind, text), ...]
                    kind ∈ {'title','time','h1','h2','lead','p','sign','table','row'}
        """
        self.doc = Document()
        self._setup_page()
        self._setup_normal()
        self._table_spec = None

        styles = self.spec['styles']
        for kind, text in blocks:
            text = '' if text is None else str(text).strip()
            if kind == 'table':
                self._finish_table()
                self._table_spec = {'row1': _split_cells(text), 'row2': None,
                                    'weights': None, 'data': []}
                continue
            if kind == 'table_header':
                self._require_table(kind)['row2'] = _split_cells(text)
                continue
            if kind == 'table_widths':
                self._require_table(kind)['weights'] = [float(value) for value in _split_cells(text)]
                continue
            if kind == 'row':
                self._require_table(kind)['data'].append(_split_cells(text))
                continue
            if kind == 'table_end':
                self._finish_table()
                continue
            self._finish_table()
            if not text:
                continue

            cfg = styles.get(kind)
            if cfg is None:
                if kind == 'time':          # 未定义 time 的格式退化为正文
                    cfg = styles['p']
                elif kind == 'sign':
                    cfg = styles.get('p', styles.get('lead'))
                else:
                    cfg = styles.get('p', styles.get('lead'))
            self._add_text(text, cfg)

        self._finish_table()
        return self.doc

    def _require_table(self, kind):
        if self._table_spec is None:
            raise ValueError(f'{kind} 必须位于 table 之后')
        return self._table_spec

    def _finish_table(self):
        if self._table_spec is None:
            return
        spec = self._table_spec
        self._table_spec = None
        row2 = spec['row2']
        if row2 is None:
            self._add_simple_table(spec['row1'], spec['data'], spec['weights'])
            return
        self._add_two_row_table(spec['row1'], row2, spec['data'], spec['weights'])

    def _add_simple_table(self, headers, data, weights):
        _validate_rows(headers, data)
        table = self.doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.style = 'Table Grid'
        _set_table_widths(table, weights, self.spec['margins'])
        for index, value in enumerate(headers):
            _set_table_cell(table.cell(0, index), value)
        for row_index, row in enumerate(data, start=1):
            for column_index, value in enumerate(row):
                _set_table_cell(table.cell(row_index, column_index), value)
            _prevent_row_split(table.rows[row_index])
        _mark_header_rows(table, 1)
        return table

    def _add_two_row_table(self, row1, row2, data, weights):
        _validate_rows(row2, data)
        _validate_group_headers(row1, len(row2))
        table = self.doc.add_table(rows=len(data) + 2, cols=len(row2))
        table.style = 'Table Grid'
        _set_table_widths(table, weights, self.spec['margins'])

        column = 0
        for raw in row1:
            text, span, vertical = _header_spec(raw)
            cell = table.cell(0, column)
            _set_table_cell(cell, text)
            if span > 1:
                cell.merge(table.cell(0, column + span - 1))
            if vertical:
                _set_vertical_merge(cell, 'restart')
            column += span

        for column, value in enumerate(row2):
            cell = table.cell(1, column)
            if value == '~':
                _set_vertical_merge(cell, None)
                _set_table_cell(cell, '')
                continue
            _set_table_cell(cell, value)

        for row_index, row in enumerate(data, start=2):
            for column_index, value in enumerate(row):
                _set_table_cell(table.cell(row_index, column_index), value)
            _prevent_row_split(table.rows[row_index])
        _mark_header_rows(table, 2)
        return table

    def save(self, output_path):
        if self.doc is None:
            raise ValueError('请先调用 build() 构建文档')
        if os.path.exists(output_path):
            raise FileExistsError(f'拒绝覆盖已有文件：{output_path}')
        self.doc.save(output_path)
        logger.info('已保存：%s', output_path)
        return output_path


def _split_cells(text):
    return [value.strip() for value in re.split(r'[；;]', text)]


def _validate_rows(headers, data):
    if not headers or any(len(row) != len(headers) for row in data):
        raise ValueError('表格数据列数必须与表头一致')


def _header_spec(value):
    span = 1
    match = re.search(r'\^(\d+)$', value)
    if match:
        span = int(match.group(1))
        value = value[:match.start()]
    vertical = value.endswith('^V')
    return (value[:-2] if vertical else value), span, vertical


def _validate_group_headers(row1, columns):
    if sum(_header_spec(value)[1] for value in row1) != columns:
        raise ValueError('第一行合并表头跨度必须等于第二行列数')


def _set_table_cell(cell, text):
    cell.text = ''
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = '仿宋'
    run.font.size = Pt(12)
    run.font.bold = False
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rfonts.set(qn(attr), '仿宋')


def _set_table_widths(table, weights, margins):
    columns = len(table.columns)
    values = weights or [1.0] * columns
    if len(values) != columns or any(value <= 0 for value in values):
        raise ValueError('表格列宽权重必须为与列数相同的正数')
    total_width = int((21.0 - margins[2] - margins[3]) / 2.54 * 1440)
    widths = [int(total_width * value / sum(values)) for value in values]
    widths[-1] = total_width - sum(widths[:-1])
    table.autofit = False
    table_width = table._tbl.tblPr.find(qn('w:tblW'))
    if table_width is None:
        table_width = OxmlElement('w:tblW')
        table._tbl.tblPr.insert(0, table_width)
    table_width.set(qn('w:type'), 'dxa')
    table_width.set(qn('w:w'), str(total_width))
    for column, width in enumerate(widths):
        table.columns[column].width = Pt(width / 20)
        for row in table.rows:
            row.cells[column].width = Pt(width / 20)


def _set_vertical_merge(cell, value):
    properties = cell._tc.get_or_add_tcPr()
    merge = OxmlElement('w:vMerge')
    if value:
        merge.set(qn('w:val'), value)
    properties.append(merge)


def _mark_header_rows(table, count):
    for row in table.rows[:count]:
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:tblHeader'))


def _prevent_row_split(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement('w:cantSplit'))


# ---------- 文书模板：六套骨架 ----------

def skeleton(fmt_name, **fields):
    """产出该格式的块序列骨架，未提供的事实写〔待补充〕"""
    def v(key, default='〔待补充〕'):
        return fields.get(key) or default

    if fmt_name in ('甲方周报', '甲方旬报'):
        tail = '周工作情况汇报' if fmt_name == '甲方周报' else '工作旬报'
        return [
            ('title', v('title', '钻试修井场信息化服务' + tail + v('period', ''))),
            ('h1', '一、本周重点工作' if fmt_name == '甲方周报' else '一、服务动态'),
            ('p', v('work')),
            ('h1', '二、本周存在问题' if fmt_name == '甲方周报' else '二、专项治理工作'),
            ('p', v('problem', '无。')),
            ('h1', '三、下周工作重点' if fmt_name == '甲方周报' else '三、存在的问题'),
            ('p', v('plan')),
        ]

    if fmt_name == '联席会汇报':
        return [
            ('title', v('title', '塔里木分院')),
            ('time', v('period')),
            ('h1', '一、周、月度会安排工作落实情况'),
            ('p', v('meeting')),
            ('h1', '二、生产运行情况'),
            ('h2', '1.工作量完成情况'),
            ('p', v('workload')),
            ('h2', '2.设备状态及动用情况'),
            ('p', v('equipment')),
            ('h2', '3.人员情况'),
            ('p', v('staff')),
            ('h2', '4.经营情况'),
            ('p', v('finance')),
            ('h1', '三、完成主要工作'),
            ('p', v('main')),
            ('h1', '四、需要科室、领导协调的问题'),
            ('p', v('issues', '无。')),
            ('h1', '五、后续的主要工作安排'),
            ('p', v('next')),
            ('sign', '汇报人：胡向峰'),
        ]

    if fmt_name == '分院党建周报':
        return [
            ('title', v('title', '党建工作及信息化服务周工作' + v('period', ''))),
            ('h1', '信息化服务'),
            ('p', v('info')),
            ('lead', '下周重点工作'),
            ('p', v('next_info')),
            ('h1', '党建工作'),
            ('p', v('party')),
        ]

    if fmt_name == '信息化周报':
        return [
            ('title', v('title', '钻试修井场信息化服务周工作情况汇报' + v('period', ''))),
            ('h1', '一、本周重点工作'),
            ('p', v('work')),
            ('h1', '二、本周专项治理工作'),
            ('p', v('special')),
            ('h1', '三、本周存在问题'),
            ('p', v('problem', '无')),
            ('h1', '四、下周工作重点'),
            ('p', v('plan')),
        ]

    if fmt_name == '个人周报':
        return [
            ('lead', '本周主要工作：'),
            ('p', v('work')),
            ('lead', '下周主要工作：'),
            ('p', v('plan')),
        ]

    raise ValueError(f'未知格式：{fmt_name}')


# ---------- 自检 ----------

def selfcheck():
    """逐套生成最小文档并回读验证核心版式。"""
    out_dir = tempfile.mkdtemp(prefix='weekly_report_selfcheck_')
    ok = True

    for name in FORMATS:
        path = os.path.join(out_dir, f'{name}.docx')
        blocks = skeleton(name, work='自检样例正文。', plan='自检样例安排。',
                          meeting='自检样例正文。', workload='自检样例工作量。',
                          equipment='自检样例设备。', staff='自检样例人员。',
                          finance='自检样例经营。', main='自检样例主要工作。',
                          next='自检样例安排。')
        if name == '联席会汇报':
            blocks.extend([
                ('table', '区块^V；周计划（天）^2'),
                ('table_header', '~；自有；合作'),
                ('table_widths', '2；1；1'),
                ('row', '塔里木；10；5'),
                ('table_end', ''),
            ])
        ReportBuilder(name).build(blocks).save(path)

        chk = Document(path)
        sec = chk.sections[0]
        t, b, l, r = FORMATS[name]['margins']
        got = (round(sec.top_margin.cm, 2), round(sec.bottom_margin.cm, 2),
               round(sec.left_margin.cm, 2), round(sec.right_margin.cm, 2))
        page_ok = got == (t, b, l, r)

        first = next((p for p in chk.paragraphs
                      if p.text.strip() and p.runs), None)
        font = size = '-'
        if first is not None:
            rf = first.runs[0]._element.rPr.rFonts
            font = (rf.get(qn('w:eastAsia')) if rf is not None
                    else FORMATS[name]['normal'][0] if FORMATS[name]['normal'] else '默认')
            size = first.runs[0].font.size.pt if first.runs[0].font.size else '默认'

        layout_ok = True
        if name == '联席会汇报':
            body = next(paragraph for paragraph in chk.paragraphs if paragraph.text == '自检样例正文。')
            spacing = body.paragraph_format.line_spacing
            line_ok = body.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY and round(spacing.pt, 1) == 28
            table = chk.tables[0]
            expected_width = int((21.0 - 2.8 - 2.6) / 2.54 * 1440)
            grid_width = sum(int(column.get(qn('w:w'))) for column in table._tbl.tblGrid.gridCol_lst)
            table_run = table.cell(1, 1).paragraphs[0].runs[0]
            table_ok = (
                len(table.rows) == 3
                and table_run.font.size.pt == 12
                and not table_run.font.bold
                and grid_width == expected_width
            )
            layout_ok = line_ok and table_ok

        status = 'OK' if page_ok and layout_ok else '版式不符'
        ok = ok and page_ok and layout_ok
        print(f'[{status:6s}] {name:8s} 边距={got} 首段字体={font} 字号={size}')

    print(f'\n输出目录：{out_dir}')
    print('自检结果：' + ('全部通过' if ok else '存在失败项'))
    return 0 if ok else 1


def load_blocks(path):
    """从文本文件读取块序列，每行 `kind|text`"""
    blocks = []
    with open(path, encoding='utf-8') as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith('#'):
                continue
            kind, _, text = line.partition('|')
            blocks.append((kind.strip(), text.strip()))
    return blocks


def main():
    logging.basicConfig(level=logging.INFO, format='%(message)s')
    ap = argparse.ArgumentParser(description='周报生成器（六套格式）')
    ap.add_argument('--list', action='store_true', help='列出全部格式')
    ap.add_argument('--selfcheck', action='store_true', help='逐套生成并验证')
    ap.add_argument('--format', dest='fmt', help='格式名')
    ap.add_argument('--blocks', help='块序列文本文件')
    ap.add_argument('--out', help='输出 docx 路径')
    args = ap.parse_args()

    if args.list:
        for name, spec in FORMATS.items():
            print(f'{name:8s} {spec["label"]}')
        return 0

    if args.selfcheck:
        return selfcheck()

    if not args.fmt:
        ap.print_help()
        return 1

    blocks = load_blocks(args.blocks) if args.blocks else skeleton(args.fmt)
    out = args.out or f'{args.fmt}.docx'
    ReportBuilder(args.fmt).build(blocks).save(out)
    print(f'已生成：{out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
