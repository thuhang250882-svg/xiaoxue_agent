import tempfile
import unittest
from pathlib import Path

from docx import Document

from annotate_docx import annotate


class AnnotateDocxTest(unittest.TestCase):
    def test_preserves_existing_comments_and_anchors_exact_text(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            output = Path(folder) / "reviewed.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("井深")
            paragraph.add_run("381mm")
            paragraph.add_run("需要复核")
            document.add_comment(paragraph.runs[:1], "已有批注", author="初审", initials="CS")
            document.save(source)

            result = annotate(
                source,
                output,
                [
                    {"match_text": "381", "comment": "381mm应写381.00mm", "author": "AI审核"},
                    {"match_text": "不存在", "comment": "不得写入"},
                ],
            )

            reopened = Document(output)
            self.assertEqual(len(reopened.comments), 2)
            self.assertEqual([comment.text for comment in reopened.comments], ["已有批注", "381mm应写381.00mm"])
            self.assertEqual([comment.author for comment in reopened.comments], ["初审", "AI审核"])
            self.assertEqual(len(result["added"]), 1)
            self.assertEqual(len(result["unmatched"]), 1)
            self.assertEqual("".join(paragraph.text for paragraph in reopened.paragraphs), "井深381mm需要复核")

    def test_rejects_overwriting_source(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            document = Document()
            document.add_paragraph("正文")
            document.save(source)
            with self.assertRaisesRegex(ValueError, "must not overwrite"):
                annotate(source, source, [])

    def test_anchors_text_inside_one_run(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            output = Path(folder) / "reviewed.docx"
            document = Document()
            document.add_paragraph("付款应当在验收后30日内完成。")
            document.save(source)

            result = annotate(
                source,
                output,
                [{"match_text": "验收后30日", "comment": "明确付款日期", "author": "AI审核"}],
            )

            reopened = Document(output)
            self.assertEqual(len(result["added"]), 1)
            self.assertEqual(len(reopened.comments), 1)
            self.assertEqual(list(reopened.comments)[0].text, "明确付款日期")
            self.assertEqual(list(reopened.comments)[0].author, "AI审核")
            self.assertEqual("".join(paragraph.text for paragraph in reopened.paragraphs), "付款应当在验收后30日内完成。")


    def _make_custom_xml_document(self, folder):
        """Body: one normal paragraph + one paragraph nested in w:customXml."""
        from docx.oxml import OxmlElement

        source = Path(folder) / "custom.docx"
        document = Document()
        document.add_paragraph("正文第一段。")
        custom = OxmlElement("w:customXml")
        nested = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "嵌在数据岛内的合同条款段落"
        run.append(text)
        nested.append(run)
        custom.append(nested)
        document.element.body.append(custom)
        document.save(source)
        return source, "嵌在数据岛内"

    def test_lifts_comment_out_of_custom_xml(self):
        with tempfile.TemporaryDirectory() as folder:
            source, match = self._make_custom_xml_document(folder)
            output = Path(folder) / "reviewed.docx"

            result = annotate(
                source,
                output,
                [{"match_text": match, "comment": "数据岛内批注必须在正文可见", "author": "AI审核"}],
            )

            self.assertEqual(len(result["added"]), 1)
            self.assertTrue(result["added"][0].get("lifted_from_custom_xml"))
            reopened = Document(output)
            self.assertEqual(len(reopened.comments), 1)
            # the annotated paragraph must now sit directly in the body flow
            body = reopened.element.body
            from docx.oxml.ns import qn

            tags = [child.tag for child in body]
            self.assertNotIn(qn("w:customXml"), tags)
            texts = [p.text for p in reopened.paragraphs]
            self.assertIn("嵌在数据岛内的合同条款段落", texts)
            # commentRangeStart must live in the lifted (visible) paragraph
            starts = body.findall(".//" + qn("w:commentRangeStart"))
            self.assertEqual(len(starts), 1)
            self.assertEqual(starts[0].getparent().tag, qn("w:p"))
            self.assertNotEqual(starts[0].getparent().getparent().tag, qn("w:customXml"))

    def test_locate_prefers_normal_flow_over_custom_xml(self):
        with tempfile.TemporaryDirectory() as folder:
            source, match = self._make_custom_xml_document(folder)
            # duplicate the same sentence into the normal body flow
            document = Document(source)
            document.add_paragraph("嵌在数据岛内的合同条款段落")
            document.save(source)
            output = Path(folder) / "reviewed.docx"

            result = annotate(
                source,
                output,
                [{"match_text": match, "comment": "应锚定到正文段落", "author": "AI审核"}],
            )

            self.assertEqual(len(result["added"]), 1)
            self.assertNotIn("lifted_from_custom_xml", result["added"][0])


if __name__ == "__main__":
    unittest.main()
